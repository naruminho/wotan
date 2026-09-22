"""Text extraction from generated artifacts - the agent's self-verification.

Used by the ``doc_read`` tool and by the verification gate (artifact evidence:
a generated PDF that extracts as empty text is a failed deliverable).
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from .csvw import read_csv
from .xlsx import read_xlsx
from .pptx import read_pptx


def read_text(path: str | Path, max_chars: int = 8000) -> dict[str, Any]:
    p = Path(path)
    suffix = p.suffix.lower().lstrip(".")
    out: dict[str, Any] = {"path": str(p), "kind": suffix or "txt"}
    if suffix == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(p))
        pages = []
        total = 0
        for i, page in enumerate(reader.pages):
            txt = (page.extract_text() or "").strip()
            total += len(txt)
            pages.append(txt[:1200])
            if total > max_chars:
                break
        out.update({"pages": len(reader.pages), "text": "\n---\n".join(pages)[:max_chars], "truncated": total > max_chars})
    elif suffix == "docx":
        import docx

        document = docx.Document(str(p))
        parts = [para.text for para in document.paragraphs if para.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        joined = "\n".join(parts)
        out.update({"text": joined[:max_chars], "truncated": len(joined) > max_chars,
                    "paragraphs": len(document.paragraphs), "tables": len(document.tables)})
    elif suffix in ("xlsx", "xlsm"):
        data = read_xlsx(p)
        lines = []
        for sheet in data["sheets"]:
            lines.append(f"# sheet: {sheet['name']}")
            lines.extend(" | ".join(r) for r in sheet["rows"])
        out.update({"text": "\n".join(lines)[:max_chars], "sheets": data["sheets"], "truncated": data.get("truncated", False)})
    elif suffix == "pptx":
        data = read_pptx(p)
        out.update({"text": "\n".join(f"[slide {s['slide']}] {s['text']}" for s in data["slides"])[:max_chars],
                    "slide_count": data["slide_count"]})
    elif suffix == "csv":
        data = read_csv(p)
        out.update({"text": "\n".join(",".join(r) for r in data["rows"])[:max_chars],
                    "rows": len(data["rows"]), "delimiter": data["delimiter"]})
    elif suffix == "json":
        text = p.read_text(encoding="utf-8", errors="replace")
        try:
            text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
        except json.JSONDecodeError:
            pass
        out.update({"text": text[:max_chars], "truncated": len(text) > max_chars})
    else:
        text = p.read_text(encoding="utf-8", errors="replace")
        out.update({"text": text[:max_chars], "truncated": len(text) > max_chars})
    return out


def sniff_header(path: str | Path) -> str:
    """Fast format sanity check without parsing: returns 'ok' | a problem note."""
    p = Path(path)
    if not p.is_file():
        return "file does not exist"
    size = p.stat().st_size
    if size == 0:
        return "file is empty (0 bytes)"
    head = p.open("rb").read(8)
    suffix = p.suffix.lower()
    if suffix == ".pdf":
        return "ok" if head.startswith(b"%PDF") else f"not a PDF (header {head[:4]!r})"
    if suffix in (".docx", ".xlsx", ".pptx"):
        return "ok" if head.startswith(b"PK") else f"not an OOXML zip (header {head[:2]!r})"
    if suffix in (".png",):
        return "ok" if head.startswith(b"\x89PNG") else f"not a PNG (header {head[:4]!r})"
    if suffix in (".jpg", ".jpeg"):
        return "ok" if head.startswith(b"\xff\xd8") else f"not a JPEG (header {head[:2]!r})"
    return "ok"
