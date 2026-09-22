"""Excel (.xlsx) writer built on openpyxl (the ``openpyxl`` extra).

Sheets are plain data described in tool calls: 2D rows with automatic typing
(numbers, ISO dates, booleans, ``=FORMULAS``), header styling, frozen header
row, autofilter and auto-fitted column widths. Also reads workbooks back
(:func:`read_xlsx`) for verification.
"""

from __future__ import annotations

import datetime as _dt
from pathlib import Path
from typing import Any


def _autotype(value: Any) -> Any:
    """Convert tool-argument strings into typed Excel cells."""
    if not isinstance(value, str):
        return value
    s = value.strip()
    if s == "":
        return None
    if s.startswith("="):
        return s  # formula
    low = s.lower()
    if low == "true":
        return True
    if low == "false":
        return False
    try:
        return int(s)
    except ValueError:
        pass
    try:
        f = float(s.replace(",", "."))
        return f
    except ValueError:
        pass
    if len(s) >= 8:
        try:
            return _dt.datetime.fromisoformat(s).replace(tzinfo=None)
        except ValueError:
            pass
    return s


def _style_sheet(ws, header_row: bool = True, freeze: bool = True, autofilter: bool = False,
                 col_widths: list[float] | None = None, currency_cols: list[int] | None = None,
                 date_cols: list[int] | None = None) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if header_row:
        fill = PatternFill("solid", fgColor="1F3A5F")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = fill
            cell.alignment = Alignment(vertical="center")
    # column widths from content
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        if col_widths and col_idx <= len(col_widths):
            ws.column_dimensions[letter].width = float(col_widths[col_idx - 1])
            continue
        best = 0
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, values_only=True):
            for v in row:
                if v is None:
                    continue
                if isinstance(v, _dt.datetime):
                    best = max(best, 11)
                else:
                    best = max(best, min(len(str(v)), 60))
        ws.column_dimensions[letter].width = max(9.0, best + 2.0)
    # number formats
    for col_idx in range(1, ws.max_column + 1):
        for cell in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for c in cell:
                if isinstance(c.value, _dt.datetime):
                    c.number_format = "DD/MM/YYYY HH:MM"
                elif isinstance(c.value, _dt.date):
                    c.number_format = "DD/MM/YYYY"
                elif isinstance(c.value, float) and currency_cols and col_idx in currency_cols:
                    c.number_format = "#,##0.00"
    if freeze and header_row:
        ws.freeze_panes = "A2"
    if autofilter and ws.max_row >= 1:
        last = get_column_letter(ws.max_column)
        ws.auto_filter.ref = f"A1:{last}{ws.max_row}"


def write_xlsx(dest: str | Path, sheets: list[dict[str, Any]], **global_opts: Any) -> dict[str, Any]:
    """sheets: [{name, rows, header?, freeze?, autofilter?, col_widths?, currency_cols?}]"""
    try:
        import openpyxl
    except ModuleNotFoundError as exc:
        raise ModuleNotFoundError(f"openpyxl is required to write .xlsx files ({exc})") from exc

    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    total_rows = 0
    for spec in sheets:
        name = str(spec.get("name") or f"Sheet{len(wb.sheetnames) + 1}")[:31]
        rows = spec.get("rows") or []
        ws = wb.create_sheet(title=name)
        for row in rows:
            typed = [_autotype(v) for v in row]
            ws.append(typed)
            total_rows += 1
        has_header = bool(rows) and isinstance(rows[0], (list, tuple))
        _style_sheet(
            ws,
            header_row=bool(spec.get("header", has_header)),
            freeze=bool(spec.get("freeze", True)),
            autofilter=bool(spec.get("autofilter", False)),
            col_widths=spec.get("col_widths"),
            currency_cols=spec.get("currency_cols"),
        )
    wb.save(str(dest))
    return {"bytes": dest.stat().st_size, "sheets": len(sheets), "rows": total_rows}


def read_xlsx(path: str | Path, max_rows: int = 100) -> dict[str, Any]:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    out: dict[str, Any] = {"sheets": []}
    for ws in wb.worksheets:
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                break
            rows.append(["" if v is None else str(v) for v in row])
        out["sheets"].append({"name": ws.title, "rows": rows, "truncated": ws.max_row > max_rows})
    wb.close()
    return out
