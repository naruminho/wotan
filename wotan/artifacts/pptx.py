"""PowerPoint (.pptx) writer built on python-pptx (the ``python-pptx`` extra).

Slides are described as data (no template lookup, so any PowerPoint/LibreOffice
opens them): title, section, bullets, two_content, image and quote layouts,
16:9 or 4:3, optional speaker notes, charts-as-images and accent shapes for a
clean, consistent design. Reads decks back via :func:`read_pptx`.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

ACCENT = (0x1F, 0x3A, 0x5F)      # deep blue
ACCENT_LIGHT = (0xE8, 0xEE, 0xF6)
TEXT = (0x2B, 0x33, 0x3D)
MUTED = (0x6B, 0x77, 0x86)
PALETTES = {
    "blue": ["1F3A5F", "2E6DA4", "5CB85C", "F0AD4E", "D9534F", "9467BD", "17A2B8", "8C564B"],
    "green": ["1E6B45", "3E9B6F", "83C69A", "F0AD4E", "D9534F", "5B8BC2", "7B6CA8", "8C6D4B"],
    "warm": ["8C4A2F", "C97B3D", "E3B23C", "7B8D42", "4E7A6A", "94516B", "5B5B8C", "8C6D4B"],
}


def _hex(color: str | tuple) -> Any:
    from pptx.util import Pt  # noqa: F401

    if isinstance(color, tuple):
        color = "".join(f"{c:02X}" for c in color)
    return color


def _add_textbox(slide, left, top, width, height, *, text="", size=18, bold=False,
                 color=TEXT, align="left", font="Calibri", italic=False, wrap=True):  # noqa: ANN001
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu, Pt

    box = slide.shapes.add_textbox(Emu(int(left)), Emu(int(top)), Emu(int(width)), Emu(int(height)))
    tf = box.text_frame
    tf.word_wrap = wrap
    lines = text.split("\n") if text else [""]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = font
        run.font.color.rgb = RGBColor.from_string(_hex(color))
    return box


def _accent_bar(slide, prs_w, top, height=0.06) -> None:  # noqa: ANN001
    from pptx.util import Emu

    shape = slide.shapes.add_shape(1, Emu(0), Emu(top), Emu(int(prs_w)), Emu(int(height * 914400)))
    from pptx.dml.color import RGBColor

    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(_hex(ACCENT))
    shape.line.fill.background()
    shape.shadow.inherit = False


def _bullets_into(tf, bullets: list, *, size=16, color=TEXT):  # noqa: ANN001
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    first = True
    for item in bullets:
        level = 0
        text = str(item)
        if text.startswith("> "):
            level = 1
            text = text[2:]
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = min(level, 4)
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size - level * 2)
        run.font.color.rgb = RGBColor.from_string(_hex(color))
        run.font.name = "Calibri"


def _add_image_slide_shape(slide, image_path: str, left, top, width, height):  # noqa: ANN001
    from PIL import Image
    from pptx.util import Emu

    img = Image.open(image_path)
    box_ratio = width / height
    img_ratio = img.width / max(img.height, 1)
    if img_ratio >= box_ratio:
        draw_w = width
        draw_h = int(width / img_ratio)
    else:
        draw_h = height
        draw_w = int(height * img_ratio)
    left_c = left + (width - draw_w) // 2
    top_c = top + (height - draw_h) // 2
    return slide.shapes.add_picture(str(image_path), Emu(int(left_c)), Emu(int(top_c)), Emu(int(draw_w)), Emu(int(draw_h)))


def write_pptx(dest: str | Path, slides: list[dict[str, Any]], *, title: str = "",
               author: str = "", aspect: str = "16:9", palette: str = "blue") -> dict[str, Any]:
    try:
        import pptx
        from pptx.util import Emu, Inches
    except ModuleNotFoundError as exc:
        raise ModuleNotFoundError(f"python-pptx is required to write .pptx files ({exc})") from exc

    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    prs = pptx.Presentation()
    if aspect == "4:3":
        prs.slide_width = Inches(10)
        prs.slide_height = Inches(7.5)
    else:
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
    W, H = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6]  # blank in the default template

    def new_slide():
        return prs.slides.add_slide(blank)

    def fill_bg(slide, hexcolor):  # noqa: ANN001
        from pptx.dml.color import RGBColor

        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(_hex(hexcolor))

    for spec in slides:
        layout = str(spec.get("layout", "bullets")).lower()
        s = new_slide()
        if layout == "title":
            fill_bg(s, ACCENT_LIGHT)
            _accent_bar(s, W, top=0.42 * H)
            t = str(spec.get("title", ""))
            _add_textbox(s, int(0.6 * 914400), int(2.1 * H / 7.5), int(W - 1.2 * 914400), int(1.6 * 914400),
                         text=t, size=40, bold=True, color=ACCENT, align="left")
            if spec.get("subtitle"):
                _add_textbox(s, int(0.62 * 914400), int(3.6 * H / 7.5), int(W - 1.2 * 914400), int(1.2 * 914400),
                             text=str(spec["subtitle"]), size=20, color=MUTED, align="left")
            if spec.get("notes"):
                s.notes_slide.notes_text_frame.text = str(spec["notes"])
        elif layout == "section":
            fill_bg(s, ACCENT)
            t = str(spec.get("title", ""))
            _add_textbox(s, int(0.8 * 914400), int(2.6 * H / 7.5), int(W - 1.6 * 914400), int(1.8 * 914400),
                         text=t, size=36, bold=True, color="FFFFFF", align="left")
            if spec.get("notes"):
                s.notes_slide.notes_text_frame.text = str(spec["notes"])
        elif layout == "image":
            if spec.get("title"):
                _add_textbox(s, int(0.6 * 914400), int(0.35 * 914400), int(W - 1.2 * 914400), int(0.9 * 914400),
                             text=str(spec["title"]), size=26, bold=True, color=ACCENT)
                _accent_bar(s, W, top=int(1.15 * 914400))
                _add_image_slide_shape(s, str(spec.get("image_path", "")), int(0.6 * 914400), int(1.45 * 914400),
                                       int(W - 1.2 * 914400), int(H - 2.1 * 914400))
            else:
                _add_image_slide_shape(s, str(spec.get("image_path", "")), int(0.4 * 914400), int(0.4 * 914400),
                                       int(W - 0.8 * 914400), int(H - 0.8 * 914400))
            if spec.get("notes"):
                s.notes_slide.notes_text_frame.text = str(spec["notes"])
        elif layout == "two_content":
            t = str(spec.get("title", ""))
            _add_textbox(s, int(0.6 * 914400), int(0.35 * 914400), int(W - 1.2 * 914400), int(0.9 * 914400),
                         text=t, size=26, bold=True, color=ACCENT)
            _accent_bar(s, W, top=int(1.15 * 914400))
            half = int(W / 2 - 0.9 * 914400)
            left_box = s.shapes.add_textbox(Emu(int(0.6 * 914400)), Emu(int(1.5 * 914400)), Emu(half), Emu(int(H - 2.0 * 914400)))
            left_box.text_frame.word_wrap = True
            _bullets_into(left_box.text_frame, list(spec.get("left") or []), size=15)
            right_box = s.shapes.add_textbox(Emu(int(W / 2 + 0.3 * 914400)), Emu(int(1.5 * 914400)), Emu(half), Emu(int(H - 2.0 * 914400)))
            right_box.text_frame.word_wrap = True
            _bullets_into(right_box.text_frame, list(spec.get("right") or []), size=15)
            if spec.get("notes"):
                s.notes_slide.notes_text_frame.text = str(spec["notes"])
        elif layout == "quote":
            fill_bg(s, "F5F7FA")
            quote = str(spec.get("text", spec.get("title", "")))
            _add_textbox(s, int(1.0 * 914400), int(2.2 * H / 7.5), int(W - 2.0 * 914400), int(2.4 * 914400),
                         text='"' + quote + '"', size=28, italic=True, color=ACCENT, align="center")
            if spec.get("author"):
                _add_textbox(s, int(1.0 * 914400), int(4.8 * H / 7.5), int(W - 2.0 * 914400), int(0.8 * 914400),
                             text="- " + str(spec["author"]), size=16, color=MUTED, align="center")
            if spec.get("notes"):
                s.notes_slide.notes_text_frame.text = str(spec["notes"])
        else:  # bullets (default)
            t = str(spec.get("title", ""))
            _add_textbox(s, int(0.6 * 914400), int(0.35 * 914400), int(W - 1.2 * 914400), int(0.9 * 914400),
                         text=t, size=26, bold=True, color=ACCENT)
            _accent_bar(s, W, top=int(1.15 * 914400))
            body_top = int(1.5 * 914400)
            body_h = int(H - body_top - 0.5 * 914400)
            bullets = list(spec.get("bullets") or [])
            image_path = spec.get("image_path")
            if image_path:
                img_w = int(W * 0.44)
                _add_image_slide_shape(s, str(image_path), int(W - img_w - 0.5 * 914400), body_top,
                                       img_w, int(body_h * 0.82))
                text_w = int(W - img_w - 1.4 * 914400)
            else:
                text_w = int(W - 1.2 * 914400)
            box = s.shapes.add_textbox(Emu(int(0.6 * 914400)), Emu(body_top), Emu(text_w), Emu(body_h))
            box.text_frame.word_wrap = True
            n = max(len(bullets), 1)
            size = 18 if n <= 6 else (15 if n <= 10 else 13)
            _bullets_into(box.text_frame, bullets, size=size)
            if spec.get("notes"):
                s.notes_slide.notes_text_frame.text = str(spec["notes"])

    prs.core_properties.title = title or Path(dest).stem
    if author:
        prs.core_properties.author = author
    prs.save(str(dest))
    return {"bytes": dest.stat().st_size, "slides": len(prs.slides.__iter__.__self__._sldIdLst) if False else len(slides)}


def read_pptx(path: str | Path, max_chars: int = 6000) -> dict[str, Any]:
    from pptx import Presentation

    prs = Presentation(str(path))
    out = []
    total = 0
    for idx, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                txt = shape.text_frame.text.strip()
                if txt:
                    texts.append(txt)
        entry = {"slide": idx, "text": " | ".join(texts)[:1200]}
        total += len(entry["text"])
        out.append(entry)
    return {"slides": out, "slide_count": len(out), "truncated": total > max_chars}
